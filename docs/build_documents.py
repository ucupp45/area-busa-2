import os
import re
import base64
import urllib.request
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

DOCS_DIR = os.path.dirname(os.path.abspath(__file__))
DIAGRAMS_DIR = os.path.join(DOCS_DIR, 'diagrams')
os.makedirs(DIAGRAMS_DIR, exist_ok=True)

# ── 1. FUNGSI RENDER MERMAID TO IMAGE ────────────────────────────
def download_mermaid(code, filename):
    filepath = os.path.join(DIAGRAMS_DIR, filename)
    if os.path.exists(filepath) and os.path.getsize(filepath) > 500:
        print(f"[CACHE] {filename} already exists.")
        return filepath
    
    try:
        # bersihkan komentar mermaid sebelum encode
        clean_code = "\n".join([line for line in code.split("\n") if not line.strip().startswith("%%")])
        b64 = base64.b64encode(clean_code.encode('utf-8')).decode('ascii')
        url = f"https://mermaid.ink/img/{b64}"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64)'})
        with urllib.request.urlopen(req, timeout=15) as resp:
            content = resp.read()
            with open(filepath, 'wb') as f:
                f.write(content)
        print(f"[SUCCESS] Downloaded {filename} ({len(content)} bytes)")
        return filepath
    except Exception as e:
        print(f"[WARN] Failed to download {filename}: {e}")
        return None

# ── 2. STYLING XML HELPERS ──────────────────────────────────────
NAVY_HEX = "0B1F44"
BLUE_HEX = "2563EB"
BG_LIGHT_HEX = "F8FAFC"
BORDER_HEX = "CBD5E1"

def set_cell_background(cell, hex_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def style_table(table, header_bg=NAVY_HEX, zebra=True):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(table.rows):
        # Prevent row split across pages
        trPr = row._element.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

        is_header = (r_idx == 0)
        if is_header:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            if is_header:
                set_cell_background(cell, header_bg)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(3)
                    for run in p.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        run.font.size = Pt(9.5)
            elif zebra and (r_idx % 2 == 1):
                set_cell_background(cell, BG_LIGHT_HEX)
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(9)
            else:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(9)

def add_callout(doc, text, title="CATATAN PENTING"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    cell = table.cell(0, 0)
    set_cell_background(cell, "EFF6FF")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

    # Set left border thick blue
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="{BLUE_HEX}"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"📌 {title}: ")
    r_title.bold = True
    r_title.font.color.rgb = RGBColor(37, 99, 235)
    r_title.font.size = Pt(9.5)
    r_text = p.add_run(text)
    r_text.font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

# ── 3. BUILD WORD DOCUMENT ───────────────────────────────────────
def generate_docx():
    doc = Document()

    # Setup page margins (1 inch)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

        # Header & Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.text = "Sistem Perencanaan Produksi & Monitoring Realisasi (Area Busa 2) | TRD, ERD & Flowchart"
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hp.runs[0].font.size = Pt(8)
        hp.runs[0].font.color.rgb = RGBColor(148, 163, 184)

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.text = "Dokumen Teknis Resmi - PT Royal Corporation (Area Busa 2)"
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.runs[0].font.size = Pt(8)
        fp.runs[0].font.color.rgb = RGBColor(148, 163, 184)

    # ── COVER / TITLE PAGE ───────────────────────────────────────
    cover_table = doc.add_table(rows=1, cols=1)
    cover_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_cell = cover_table.cell(0, 0)
    set_cell_background(c_cell, NAVY_HEX)
    set_cell_margins(c_cell, top=360, bottom=360, left=360, right=360)
    
    cp = c_cell.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_pre = cp.add_run("STANDAR OPERASIONAL & ARSITEKTUR TEKNIS\n")
    r_pre.font.size = Pt(11)
    r_pre.font.bold = True
    r_pre.font.color.rgb = RGBColor(168, 216, 240) # Accent blue

    r_main = cp.add_run("SISTEM PERENCANAAN PRODUKSI & MONITORING REALISASI\n")
    r_main.font.size = Pt(18)
    r_main.font.bold = True
    r_main.font.color.rgb = RGBColor(255, 255, 255)

    r_sub = cp.add_run("Lini Pemotongan Busa: Area Busa 2\n(Mesin Horizontal, Vertikal, Carousel)\n")
    r_sub.font.size = Pt(13)
    r_sub.font.color.rgb = RGBColor(226, 232, 240)

    doc.add_paragraph().paragraph_format.space_before = Pt(18)

    # Meta box
    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Nama Dokumen", "Technical Requirements Document (TRD), ERD & Flowchart"),
        ("Versi Dokumen", "1.0 (Production-Ready)"),
        ("Tanggal Terbit", "16 Oktober 2026"),
        ("Target Pengguna", "PPC, Operator Mesin, Supervisor Produksi, Superadmin"),
        ("Status Pengesahan", "Approved / Resmi")
    ]
    for idx, (label, val) in enumerate(meta_data):
        row = meta_table.rows[idx]
        row.cells[0].paragraphs[0].add_run(label).bold = True
        row.cells[0].paragraphs[0].runs[0].font.size = Pt(9.5)
        row.cells[1].paragraphs[0].add_run(val)
        row.cells[1].paragraphs[0].runs[0].font.size = Pt(9.5)
    style_table(meta_table, header_bg="1E293B", zebra=True)

    doc.add_page_break()

    # ── 1. TECHNICAL REQUIREMENTS DOCUMENT (TRD) ─────────────────
    h1 = doc.add_heading("1. TECHNICAL REQUIREMENTS DOCUMENT (TRD)", level=1)
    h1.runs[0].font.color.rgb = RGBColor(11, 31, 68)

    # 1.1 Latar Belakang & Tujuan
    h2 = doc.add_heading("1.1 Latar Belakang & Tujuan", level=2)
    h2.runs[0].font.color.rgb = RGBColor(37, 99, 235)
    doc.add_paragraph(
        "Lini produksi Area Busa 2 bertanggung jawab dalam pemrosesan balokan busa (polyurethane foam blocks) "
        "menjadi produk jadi maupun setengah jadi sesuai spesifikasi Work Order (WO). Lini ini mengoperasikan "
        "tiga jenis mesin pemotongan utama: Mesin Horizontal (HRZ), Mesin Vertikal (VRT), dan Mesin Carousel (CRS)."
    )
    doc.add_paragraph(
        "Sistem ini dirancang untuk menyelesaikan tantangan operasional manual:\n"
        "1. Perhitungan Otomatis Kubikasi & Cycle Time: Meniadakan kalkulasi rumus spreadsheet manual untuk menentukan kebutuhan man power dan jam mesin per shift.\n"
        "2. Sinkronisasi Real-Time Target vs Aktual: Menghubungkan rencana PPC secara langsung ke tablet/layar operator di lini mesin.\n"
        "3. Pengukuran Efektivitas Mesin (OEE): Menghitung Availability, Performance, dan Quality Rate secara otomatis dari data transaksi.\n"
        "4. Pengendalian Limbah Sisa Potong: Mencatat sisa potongan blok busa ke dalam inventori terpusat agar dapat dialokasikan pada pesanan berikutnya."
    )

    # 1.2 Arsitektur Sistem
    h2 = doc.add_heading("1.2 Arsitektur Sistem (Three-Tier Architecture)", level=2)
    h2.runs[0].font.color.rgb = RGBColor(37, 99, 235)
    doc.add_paragraph(
        "Sistem dibangun menggunakan arsitektur Three-Tier Client-Server yang memisahkan presentasi antarmuka, "
        "logika bisnis, dan persistensi basis data:"
    )

    arch_table = doc.add_table(rows=4, cols=3)
    arch_data = [
        ("Layer Arsitektur", "Teknologi & Framework", "Deskripsi Peran & Tanggung Jawab"),
        ("Presentation Layer", "React 19, Vite 8, Tailwind CSS v4, Recharts, Lucide Icons", "Antarmuka Single Page Application (SPA) responsif, dark mode toggle, visualisasi grafik OEE, serta validasi form interaktif."),
        ("Application Layer", "Node.js (v18+ LTS), Express.js 5.x, Multer, xlsx (SheetJS)", "RESTful API Server, Engine Parsing SPK, Engine Kalkulasi OEE & Kubikasi, Server-Side Role Guard, Controller Stok Sisa Potong."),
        ("Data Layer", "MySQL 8.0+ Database (dashboard_vertikal), InnoDB Engine", "Penyimpanan relasional terindeks ACID-compliant, connection pool mysql2 teroptimasi, riwayat transaksi lengkap.")
    ]
    for r_idx, row in enumerate(arch_table.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.paragraphs[0].text = arch_data[r_idx][c_idx]
    style_table(arch_table)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 1.3 Tech Stack & Spesifikasi Lingkungan
    h2 = doc.add_heading("1.3 Spesifikasi Tech Stack", level=2)
    h2.runs[0].font.color.rgb = RGBColor(37, 99, 235)

    stack_table = doc.add_table(rows=6, cols=3)
    stack_data = [
        ("Komponen", "Versi", "Kegunaan Utama"),
        ("React", "19.2", "Pembangunan UI reaktif berbasis komponen modular."),
        ("Vite", "8.2", "Build tool & development server super cepat dengan HMR."),
        ("Tailwind CSS", "4.3", "Styling utility-first dengan sistem tema warna modern."),
        ("Node.js / Express", "18+ / 5.2", "Runtime asynchronous backend RESTful API."),
        ("MySQL Database", "8.0+", "RDBMS dengan engine InnoDB dan indexing B-Tree.")
    ]
    for r_idx, row in enumerate(stack_table.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.paragraphs[0].text = stack_data[r_idx][c_idx]
    style_table(stack_table)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 1.4 Kebutuhan Fungsional (FR)
    h2 = doc.add_heading("1.4 Kebutuhan Fungsional (Functional Requirements)", level=2)
    h2.runs[0].font.color.rgb = RGBColor(37, 99, 235)

    fr_items = [
        ("FR-01: Autentikasi & Role-Based Access Control (RBAC)", 
         "Sistem menerapkan kontrol akses berbasis peran (Superadmin, PPC, Operator). Header 'x-user-role' diverifikasi pada setiap pemanggilan API sensitif."),
        ("FR-02: Modul Perencanaan SPK (Planning PPC)", 
         "PPC dapat mengunggah file Excel SPK (.xls/.xlsx) atau menempel teks tabel spreadsheet. Sistem otomatis mem-parsing Work Order, mencocokkan cycle time matrix, dan menghitung total kubikasi (M3) serta man power ideal."),
        ("FR-03: Modul Realisasi Pemotongan Mesin (Operator)", 
         "Operator mencatat hasil pemotongan balok aktual (Vertikal & Horizontal), waste kulit, waste samping, waste produksi (defect), serta kendala waktu downtime mesin."),
        ("FR-04: Modul Manajemen Stok Sisa Potong (Sispot)", 
         "Sisa pemotongan yang masih bernilai otomatis tersimpan ke inventori sispot. Terdapat analitik status (Tersedia/Terpakai), Top 5 Kualitas, dan pengelompokan Aging (Hari ini, 1-3 hari, 4-7 hari, >7 hari). Operasi Create, Update, dan Delete (CRUD) dilindungi ketat dan hanya dapat dijalankan oleh Superadmin."),
        ("FR-05: Modul OEE Metric Engine", 
         "Menghitung metrik Overall Equipment Effectiveness secara otomatis: Availability = (Operating Time / Planned Time), Performance = (Actual Output / Ideal Output), Quality = (Good Output / Total Output). Grafik tren 7 hari dilengkapi garis ambang batas 80%.")
    ]
    for title, desc in fr_items:
        p = doc.add_paragraph()
        r = p.add_run(f"• {title}: ")
        r.bold = True
        r.font.color.rgb = RGBColor(15, 23, 42)
        p.add_run(desc)

    add_callout(doc, "Seluruh operasi penambahan manual, pengubahan ukuran, dan penghapusan data Stok Sisa Potong (Sispot) memiliki proteksi lapis ganda di antarmuka frontend dan validasi server-side HTTP 403 Forbidden.", "KEAMANAN AKSES CRUD SUPERADMIN")

    # 1.5 Kebutuhan Non-Fungsional (NFR)
    h2 = doc.add_heading("1.5 Kebutuhan Non-Fungsional (Non-Functional Requirements)", level=2)
    h2.runs[0].font.color.rgb = RGBColor(37, 99, 235)

    nfr_table = doc.add_table(rows=5, cols=2)
    nfr_data = [
        ("Parameter NFR", "Standar & Target Spesifikasi"),
        ("Waktu Respon (Performance)", "Rata-rata respon API <= 200 ms. Waktu render dashboard komparasi <= 1.5 detik."),
        ("Kapasitas Penyimpanan (Scalability)", "Mampu menampung >= 1.000.000 transaksi pemotongan dengan pengindeksan B-Tree pada kolom no_barcode, tanggal, dan identitas_id."),
        ("Ketersediaan (Availability)", "Toleransi uptime operasional 99.5% selama 3 shift kerja pabrik (24 jam operasional)."),
        ("Integritas Data (Reliability)", "Menerapkan transaksi MySQL dan cascading foreign key untuk menjaga konsistensi identitas balok dan rincian pemotongan.")
    ]
    for r_idx, row in enumerate(nfr_table.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.paragraphs[0].text = nfr_data[r_idx][c_idx]
    style_table(nfr_table)

    doc.add_page_break()

    # ── 2. ENTITY RELATIONSHIP DIAGRAM (ERD) ──────────────────────
    h1 = doc.add_heading("2. ENTITY RELATIONSHIP DIAGRAM (ERD)", level=1)
    h1.runs[0].font.color.rgb = RGBColor(11, 31, 68)

    doc.add_paragraph(
        "Basis data relational 'dashboard_vertikal' menggunakan engine InnoDB dengan relasi referensial "
        "terstruktur untuk mendukung siklus hidup pemotongan busa secara menyeluruh:"
    )

    # Diagram ERD
    erd_code = """erDiagram
    UPLOAD_HISTORY ||--o{ SPK_ITEMS : "memiliki detail"
    REALISASI_IDENTITAS ||--o{ REALISASI_VERTICAL : "potong vertikal"
    REALISASI_IDENTITAS ||--o{ REALISASI_CAROUSEL : "potong horizontal"
    REALISASI_IDENTITAS ||--o{ STOK_SISA_POTONG : "menghasilkan sispot"
    MACHINES ||--o{ MACHINE_HEALTH : "status mesin"

    UPLOAD_HISTORY {
        int id PK
        varchar nama_file
        datetime tanggal_upload
        int shift
        int total_item
    }
    SPK_ITEMS {
        int id PK
        int history_id FK
        text no_wo
        varchar nama_item
        decimal qty
        decimal estimasi_waktu_menit
        decimal ukuran_p
        decimal ukuran_l
        decimal ukuran_t
    }
    REALISASI_IDENTITAS {
        int id PK
        varchar no_barcode
        varchar kualitas
        date tanggal
        varchar shift
        decimal asal_p
        decimal asal_l
        decimal asal_t
        int asal_qty
    }
    REALISASI_VERTICAL {
        int id PK
        int identitas_id FK
        decimal hasil_p
        decimal hasil_l
        decimal hasil_t
        int hasil_qty
        varchar no_mesin
    }
    REALISASI_CAROUSEL {
        int id PK
        int identitas_id FK
        varchar no_wo
        decimal hasil_p
        decimal hasil_l
        decimal hasil_t
        int hasil_qty
        varchar no_mesin
    }
    STOK_SISA_POTONG {
        int id PK
        int identitas_id FK
        varchar no_barcode
        varchar kualitas
        varchar sumber_potong
        float p
        float l
        float t
        int qty
        tinyint is_used
    }
    """
    erd_img = download_mermaid(erd_code, "erd_diagram.png")
    if erd_img:
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(erd_img, width=Inches(6.2))
        
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap = cap.add_run("Gambar 2.1: Entity Relationship Diagram Basis Data Area Busa 2")
        r_cap.font.italic = True
        r_cap.font.size = Pt(8.5)
        r_cap.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_heading("2.1 Kamus Data Entitas Utama", level=2).runs[0].font.color.rgb = RGBColor(37, 99, 235)

    dict_table = doc.add_table(rows=7, cols=4)
    dict_data = [
        ("Nama Tabel", "Primary Key", "Foreign Key", "Fungsi Bisnis"),
        ("upload_history", "id", "-", "Header sesi upload file SPK Excel harian oleh PPC."),
        ("spk_items", "id", "history_id -> upload_history.id", "Detail baris pesanan Work Order hasil parsing."),
        ("realisasi_identitas", "id", "-", "Header balokan busa asal (barcode, dimensi, operator, shift)."),
        ("realisasi_vertical", "id", "identitas_id -> realisasi_identitas.id", "Detail hasil potong, waste, dan downtime mesin Vertikal."),
        ("realisasi_carousel", "id", "identitas_id -> realisasi_identitas.id", "Detail hasil potong, waste, dan downtime mesin Horizontal/Carousel."),
        ("stok_sisa_potong", "id", "identitas_id -> realisasi_identitas.id", "Inventori sisa potongan busa (auto-generated & manual Superadmin).")
    ]
    for r_idx, row in enumerate(dict_table.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.paragraphs[0].text = dict_data[r_idx][c_idx]
    style_table(dict_table)

    doc.add_page_break()

    # ── 3. FLOWCHART SISTEM ───────────────────────────────────────
    h1 = doc.add_heading("3. FLOWCHART & ALUR PROSES SISTEM", level=1)
    h1.runs[0].font.color.rgb = RGBColor(11, 31, 68)

    # 3.1 Flowchart 1
    h2 = doc.add_heading("3.1 Flowchart Global Sistem Terintegrasi", level=2)
    h2.runs[0].font.color.rgb = RGBColor(37, 99, 235)
    doc.add_paragraph(
        "Diagram alir global menggambarkan integrasi lintas departemen antara PPC (Perencanaan), "
        "Operator Produksi (Eksekusi Mesin), dan Manajemen Pabrik (Evaluasi Performa & OEE):"
    )

    fc1_code = """flowchart TD
    Start([Mulai]) --> Login[Pengguna Login ke Sistem]
    Login --> RoleCheck{Cek Role Pengguna}
    
    RoleCheck -->|PPC / Superadmin| MenuPlanning[Buka Menu Planning]
    MenuPlanning --> InputSPK[Upload SPK Excel / Copas Text]
    InputSPK --> ParseEngine[Engine Parsing & Auto-Fill Kubikasi]
    ParseEngine --> SimpanSPK[(Database: upload_history & spk_items)]
    SimpanSPK --> DashPlan[Pantau Dashboard Planning & Prediksi OEE]

    RoleCheck -->|Operator / Superadmin| MenuRealisasi[Buka Menu Realisasi]
    MenuRealisasi --> PilihMesin{Pilih Mesin: Vertikal / Horizontal}
    PilihMesin --> FormPotong[Input Potongan Aktual & Waste]
    FormPotong --> CekSispot{Ada Sisa Potong?}
    CekSispot -->|Ya| AutoSispot[Generate Otomatis Stok Sisa Potong]
    CekSispot -->|Tidak| SimpanReal[(Database: realisasi_identitas & detail)]
    AutoSispot --> SimpanReal

    SimpanSPK --> CompareEngine[Engine Komparasi Plan vs Real]
    SimpanReal --> CompareEngine
    CompareEngine --> DashComp[Dashboard Planning vs Realisasi]
    DashComp --> AnalisisOEE[Analisis OEE, Downtime & Gap Kubikasi]
    AnalisisOEE --> End([Selesai])
    """
    fc1_img = download_mermaid(fc1_code, "flowchart1_global.png")
    if fc1_img:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(fc1_img, width=Inches(5.5))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.add_run("Gambar 3.1: Flowchart Alur Kerja Sistem Terintegrasi Area Busa 2").font.italic = True

    doc.add_page_break()

    # 3.2 Flowchart 2 (PPC Planning)
    h2 = doc.add_heading("3.2 Flowchart Perencanaan SPK (PPC Planning)", level=2)
    h2.runs[0].font.color.rgb = RGBColor(37, 99, 235)
    doc.add_paragraph("Alur kerja pengunggahan, ekstraksi data, dan perhitungan analitik rencana produksi harian:")

    fc2_code = """flowchart TD
    A[PPC Buka Form Insert SPK] --> B{Pilih Metode Input}
    B -->|Upload File| C1[Pilih File Excel .xls / .xlsx]
    C1 --> C2[Parser XLSX SheetJS]
    B -->|Copy-Paste| D1[Paste Teks Tabel Spreadsheet]
    D1 --> D2[Regex & Tab Delimiter Parser]
    
    C2 --> E[Ekstraksi Baris WO & Dimensi P x L x T]
    D2 --> E
    E --> F[Pencocokan Nama Kualitas ke cycle_time_matrix]
    F --> G[Hitung Kubikasi M3 = P x L x T x Qty / 1.000.000]
    G --> H[Kalkulasi Estimasi Waktu Siklus & Man Power]
    H --> I[Tampilkan Tabel Pratinjau Rekapitulasi]
    I --> J{Konfirmasi Simpan?}
    J -->|Ya| K[POST /api/spk/upload ke Database]
    J -->|Batal| I
    K --> L[Tampil di Dashboard Planning & OEE Rencana]
    """
    fc2_img = download_mermaid(fc2_code, "flowchart2_planning.png")
    if fc2_img:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(fc2_img, width=Inches(5.2))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.add_run("Gambar 3.2: Flowchart Alur Perencanaan SPK oleh PPC").font.italic = True

    # 3.3 Flowchart 3 (Operator Cutting)
    h2 = doc.add_heading("3.3 Flowchart Realisasi Pemotongan (Operator Mesin)", level=2)
    h2.runs[0].font.color.rgb = RGBColor(37, 99, 235)
    doc.add_paragraph("Alur kerja operator dalam mengeksekusi pemotongan balok, input waste, dan pelaporan kendala:")

    fc3_code = """flowchart TD
    A[Operator Buka Form Input Realisasi] --> B[Pilih Shift, Tanggal & Operator]
    B --> C[Input / Scan Barcode Balokan Asal]
    C --> D{Pilih Jenis Mesin}
    D -->|Vertikal| V1[Input Dimensi Jadi & Waste Samping]
    D -->|Horizontal/Carousel| H1[Input Dimensi Jadi & Waste Kulit]
    V1 --> S1{Ada Sisa Potong?}
    H1 --> S1
    S1 -->|Ya| S2[Input Dimensi Sisa Potong P x L x T & Qty]
    S1 -->|Tidak| DT1{Ada Kendala Mesin?}
    S2 --> DT1
    DT1 -->|Ya| DT2[Input Kode Downtime & Jam Awal-Akhir]
    DT1 -->|Tidak| SAVE[Klik Simpan Realisasi]
    DT2 --> SAVE
    SAVE --> API[Kirim POST ke Backend API]
    API --> DB[(Database: realisasi_identitas & stok_sisa_potong)]
    DB --> OUT[Update Stok Sispot & Muncul di Log Hasil]
    """
    fc3_img = download_mermaid(fc3_code, "flowchart3_realisasi.png")
    if fc3_img:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(fc3_img, width=Inches(5.2))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.add_run("Gambar 3.3: Flowchart Alur Eksekusi Pemotongan Operator").font.italic = True

    doc.add_page_break()

    # 3.4 Flowchart 4 (Stok Sisa Potong Lifecycle)
    h2 = doc.add_heading("3.4 Flowchart Siklus Hidup Stok Sisa Potong (Sispot)", level=2)
    h2.runs[0].font.color.rgb = RGBColor(37, 99, 235)
    doc.add_paragraph("Alur tata kelola inventori sisa potong dengan hak akses khusus Superadmin:")

    fc4_code = """flowchart TD
    S1[Hasil Potong Operator] -->|Otomatis Masuk| INV[(Inventori Stok Sisa Potong)]
    S2[Temuan Busa atau Input Manual] --> S3[Superadmin Buka Form Tambah]
    S3 --> S4{Cek Role Superadmin}
    S4 -->|Bukan Superadmin| S5[Akses Ditolak: 403 Forbidden]
    S4 -->|Superadmin Valid| S6[Input Dimensi P x L x T & Qty]
    S6 --> INV
    
    INV --> M1[Sistem Hitung Aging Hari Simpan]
    M1 --> M2[Kategori: Hari Ini, 1-3 Hari, 4-7 Hari, gt 7 Hari]
    M2 --> USE{Sispot Dipakai untuk WO Baru?}
    USE -->|Ya| U1[Update Status: Terpakai]
    USE -->|Batal atau Rusak| U2{Superadmin Mau Hapus?}
    U2 -->|Ya| U3[Aksi Hapus Superadmin]
    U3 --> DEL[(Hapus dari Database)]
    """
    fc4_img = download_mermaid(fc4_code, "flowchart4_sispot.png")
    if fc4_img:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(fc4_img, width=Inches(5.0))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.add_run("Gambar 3.4: Flowchart Alur Inventori Stok Sisa Potong").font.italic = True

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # ── 4. LAMPIRAN PENGESAHAN ───────────────────────────────────
    h1 = doc.add_heading("4. LEMBAR PENGESAHAN DOKUMEN", level=1)
    h1.runs[0].font.color.rgb = RGBColor(11, 31, 68)

    doc.add_paragraph(
        "Dokumen spesifikasi teknis dan diagram alir ini telah ditinjau dan dinyatakan sah "
        "sebagai acuan standar operasional sistem pada lini Area Busa 2."
    )

    sign_table = doc.add_table(rows=3, cols=3)
    sign_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sign_headers = [("Disusun Oleh:", "Diperiksa Oleh:", "Disetujui Oleh:")]
    sign_table.rows[0].cells[0].paragraphs[0].text = "Disusun Oleh (Pengembang):"
    sign_table.rows[0].cells[1].paragraphs[0].text = "Diperiksa Oleh (PPC Lead):"
    sign_table.rows[0].cells[2].paragraphs[0].text = "Disetujui Oleh (Factory Head):"
    
    sign_table.rows[1].cells[0].paragraphs[0].text = "\n\n\n_________________________\nTim Pengembang Sistem"
    sign_table.rows[1].cells[1].paragraphs[0].text = "\n\n\n_________________________\nSupervisor PPC"
    sign_table.rows[1].cells[2].paragraphs[0].text = "\n\n\n_________________________\nKepala Pabrik / Produksi"

    sign_table.rows[2].cells[0].paragraphs[0].text = "Tanggal: 16 Oktober 2026"
    sign_table.rows[2].cells[1].paragraphs[0].text = "Tanggal: 16 Oktober 2026"
    sign_table.rows[2].cells[2].paragraphs[0].text = "Tanggal: 16 Oktober 2026"
    style_table(sign_table, header_bg="1E293B", zebra=False)

    docx_path = os.path.join(DOCS_DIR, "TRD_ERD_FLOWCHART.docx")
    doc.save(docx_path)
    print(f"\n[SUCCESS] Word Document created: {docx_path}")
    return docx_path

if __name__ == "__main__":
    docx_path = generate_docx()
    pdf_path = os.path.join(DOCS_DIR, "TRD_ERD_FLOWCHART.pdf")
    print(f"\n[INFO] Converting {docx_path} to PDF using docx2pdf...")
    try:
        from docx2pdf import convert
        convert(docx_path, pdf_path)
        print(f"[SUCCESS] PDF Document created: {pdf_path}")
    except Exception as e:
        print(f"[ERROR] Failed to convert to PDF via docx2pdf: {e}")
        # fallback using win32com if needed
        try:
            import win32com.client
            word = win32com.client.Dispatch('Word.Application')
            doc = word.Documents.Open(docx_path)
            doc.SaveAs(pdf_path, FileFormat=17) # 17 = wdFormatPDF
            doc.Close()
            word.Quit()
            print(f"[SUCCESS] PDF Document created via win32com: {pdf_path}")
        except Exception as e2:
            print(f"[ERROR] win32com fallback also failed: {e2}")
